from pathlib import Path
from shutil import copy2

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).parent
TUTORIALS = ROOT / 'Tutorials'
CURRICULUM = Path(r'C:\Users\Rob011235\Dropbox\CNM\_Curriculum\CIST 2220 C# .NET I\MAUI')
VISUALS = {
    8: {
        1: [('module-8-concept.png', 'The toolkit generates binding and command plumbing while the view model keeps application rules.')],
        2: [('editor-toolkit.png', 'The editor still shows a live total and enabled Submit after its properties become generated.')],
        5: [('bom-toolkit.png', 'Generated commands still add and change Bill of Materials rows.')],
    },
    9: {
        2: [('module-9-concept.png', 'The view model asks for an edit or confirmation through an interface; the MAUI adapter opens pages and alerts.')],
        3: [('injected-editor.png', 'The injected editor opens with the selected item ready to edit.')],
        4: [('confirmation-service.png', 'The interaction service owns the platform confirmation alert; the view model uses its result.')],
    },
    10: {
        2: [('module-10-concept.png', 'The view model uses the service contract; the JSON implementation saves and reloads the collection.')],
        4: [('saved-item.png', 'The new item appears after it has been saved.'),
            ('reloaded-item.png', 'The same item appears after closing and reopening the application.')],
    },
    11: {
        2: [('module-11-concept.png', 'Plan views, view models, and services before implementing the redo.')],
        6: [('reference-editor.png', 'The completed editor binds its fields and commands to an editor view model.')],
        8: [('reference-bom.png', 'The completed Bill of Materials shows one row and the Add, Edit, and Delete actions.')],
        9: [('reference-home.png', 'The completed Home page is reached through Shell navigation.')],
    },
}


def insert_after(anchor, paragraph):
    anchor._p.addnext(paragraph._p)
    return paragraph


for module, visual_steps in VISUALS.items():
    folder = TUTORIALS / f'Module {module}'
    path = next(folder.glob('*.docx'))
    doc = Document(path)
    if doc.inline_shapes:
        raise RuntimeError(f'{path} already contains images; avoid duplicate insertion')

    for step, visuals in visual_steps.items():
        heading = next(
            (i for i, p in enumerate(doc.paragraphs)
             if p.text.startswith(f'Step {step} ') and p.style.name.startswith('Heading')),
            None,
        )
        if heading is None:
            raise RuntimeError(f'Step {step} not found in {path}')
        following = doc.paragraphs[heading + 1:]
        check = next((p for p in following if p.text.startswith('Check. ')), None)
        if check is None:
            raise RuntimeError(f'Check paragraph for step {step} not found in {path}')
        anchor = check
        for filename, caption in visuals:
            image_path = folder / filename
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.keep_with_next = True
            picture.add_run().add_picture(str(image_path), width=Inches(6.7))
            anchor = insert_after(anchor, picture)
            label = doc.add_paragraph('Figure. ' + caption, style='Caption')
            label.paragraph_format.keep_together = True
            anchor = insert_after(anchor, label)

    doc.save(path)
    target = CURRICULUM / f'Module {module}' / path.name
    copy2(path, target)
    print(module, path, len(Document(path).inline_shapes), target)
