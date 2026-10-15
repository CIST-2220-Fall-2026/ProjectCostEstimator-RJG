from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
output = root / 'Tutorials' / 'MAUI_IoT_Project_Cost_Estimator_11_Module_Arc.docx'
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(.7)
section.left_margin = section.right_margin = Inches(.85)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 20), ('Heading 1', 14), ('Heading 2', 11.5)]:
    styles[name].font.name = 'Aptos'
    styles[name].font.size = Pt(size)
    styles[name].font.color.rgb = RGBColor(0, 0, 0)
ppr = styles['Title'].element.get_or_add_pPr()
borders = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'nil')
borders.append(bottom)
ppr.append(borders)

doc.add_paragraph('IoT Project Cost Estimator Eleven Module Learning Arc', 'Title')
doc.add_paragraph('Students build one working .NET MAUI estimator while moving from named controls and click handlers to a production-style MVVM design. Each of the first ten modules changes one responsibility in the same app. Module 11 starts again from a clean template so students assemble the final design independently.')
doc.add_heading('How the sequence teaches MVVM', 1)
doc.add_paragraph('Modules 1 and 2 establish the feature workflow. Modules 3 through 7 reveal binding, commands, collection notifications, view models, services, and dependency injection one concept at a time. Modules 8 through 10 replace handwritten plumbing with the MVVM Toolkit, clarify the UI boundary, and add persistence and tests. The redo lab asks students to apply those choices from the start.')

modules = [
    (1, 'Model and Item Editor', 'Collect item values with named controls, calculate cost, and validate input.', 'A valid item is created; invalid input does not crash the app.'),
    (2, 'Bill of Materials workflow', 'Add, edit, delete, and cancel using pages, navigation, a list, and manual refresh.', 'Every BOM action and Help navigation works.'),
    (3, 'Editor binding', 'Introduce BindingContext, two-way Entry binding, INotifyPropertyChanged, and a live total.', 'The editor previews cost while typing.'),
    (4, 'Commands and validation', 'Replace editor button handlers with commands; connect validity to Submit availability.', 'Submit enables only for valid values.'),
    (5, 'BOM view model', 'Move list actions to a view model and pass row items with CommandParameter.', 'Row commands work while refresh remains explicit.'),
    (6, 'Observable collection', 'Replace manual list refresh with ObservableCollection notifications.', 'Rows update on Add, Replace, and Remove.'),
    (7, 'Services and dependency injection', 'Move BOM data to IBomService and inject service, view model, and page.', 'The same workflow runs with clear data ownership.'),
    (8, 'MVVM Toolkit refactor', 'Use generated observable properties and commands while retaining explicit validation rules.', 'The app behaves as before with less notification and command boilerplate.'),
    (9, 'Navigation and dialog boundaries', 'Inject the editor view model and use a small interaction interface for editor results and delete confirmation.', 'View model decisions can run without a MAUI page.'),
    (10, 'Persistence and tests', 'Save the BOM as JSON, load it on return, surface errors, and test services and view models.', 'Changes survive restart and automated behavior checks pass.'),
    (11, 'MVVM redo from the start', 'Rebuild the full app from a clean MAUI starter using the final architecture.', 'The independent reference builds; the redo passes UI and automated checks.'),
]

doc.add_heading('Module sequence', 1)
for number, title, action, check in modules:
    doc.add_heading(f'Module {number} {title}', 2)
    doc.add_paragraph(action)
    p = doc.add_paragraph()
    p.add_run('Checkpoint. ').bold = True
    p.add_run(check)

doc.add_heading('Teaching and assessment boundaries', 1)
doc.add_paragraph('Keep a working app at the end of every module. Build the Windows target and check Add, Edit, Cancel, and Delete whenever those paths change. Modules 8 and 9 are refactors: the visible behavior should stay familiar. Module 10 deliberately changes storage from memory to disk. Module 11 begins with the clean template and should be assessed as an independent rebuild, using the reference only after each attempt.')
doc.add_paragraph('Navigation and platform alerts remain UI concerns, implemented through a small MAUI adapter. The view models own decisions and depend on interfaces. Production-style MVVM here means deliberate boundaries, generated notification and command code, persistence, and tests; it does not require empty page code files.')
doc.add_heading('Deliverables', 1)
doc.add_paragraph('Modules 1 through 11 each have a Word tutorial and captured app screens in the curriculum folder. The Module 11 branch also contains Module11Starter, Module11Reference, and Module11Reference.Tests. Feature branches for Modules 8 through 11 are pushed separately so each lesson can be reviewed at its own endpoint.')
doc.save(output)
curriculum = Path(r'C:\Users\Rob011235\Dropbox\CNM\_Curriculum\CIST 2220 C# .NET I\MAUI') / output.name
curriculum.write_bytes(output.read_bytes())
print(output)
print(curriculum)
