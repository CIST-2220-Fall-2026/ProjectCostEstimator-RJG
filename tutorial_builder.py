from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT / 'Tutorials'
OUT.mkdir(exist_ok=True)

def build(number, title, introduction, steps, files, references, starting_point=None):
    folder = OUT / f'Module {number}'
    folder.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.85)
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 21), ('Heading 1', 15), ('Heading 2', 12)]:
        styles[name].font.name = 'Aptos Display' if name == 'Title' else 'Aptos'
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = __import__('docx').shared.RGBColor(0, 0, 0)
    title_ppr = styles['Title'].element.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'nil')
    borders.append(bottom)
    title_ppr.append(borders)
    doc.add_paragraph(f'Module {number} {title}', 'Title')
    doc.add_paragraph(introduction)
    doc.add_heading('Starting point and learning goals', 1)
    doc.add_paragraph(starting_point or f'Start from the completed Module {number-1} branch. Create a new feature branch before editing. Build the Windows target at each checkpoint.')
    doc.add_paragraph('Learning goals: ' + '; '.join(s[1] for s in steps) + '.')
    for i, (heading, goal, work, check) in enumerate(steps, 1):
        doc.add_heading(f'Step {i} {heading}', 1)
        doc.add_paragraph('Concept. ' + goal)
        doc.add_paragraph('Do. ' + work)
        doc.add_paragraph('Check. ' + check)
    doc.add_heading('Further reading', 1)
    for ref in references:
        doc.add_paragraph(ref)
    doc.add_heading('Source checkpoints', 1)
    doc.add_paragraph('The files below show key parts of the completed module. The reference project contains every file. Read these after following the small steps above, or use them to compare your result.')
    for label, path in files:
        doc.add_heading(label, 2)
        if not path.exists():
            doc.add_paragraph('This file is supplied by the previous module.')
            continue
        for line in path.read_text(encoding='utf-8').splitlines():
            p = doc.add_paragraph(style='Normal')
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.left_indent = Inches(.15)
            run = p.add_run(line or ' ')
            run.font.name = 'Consolas'
            run.font.size = Pt(7.75)
    output = folder / f'Module_{number}_{title.replace(" ", "_")}_Tutorial.docx'
    doc.save(output)
    print(output)
    return output

if __name__ == '__main__':
    import sys
    number = int(sys.argv[1])
    source = ROOT / 'ProjectCostEstimator'
    if number == 8:
        build(8, 'MVVM Toolkit Refactor',
            'Refactor the working estimator to use generated observable properties and commands. The app keeps the same behavior while the repeated notification and command setup code disappears. The validation rules remain explicit so you can distinguish framework plumbing from application decisions.',
            [
                ('Add the package', 'Identify which code the toolkit can generate', 'Add CommunityToolkit.Mvvm version 8.4.2 to the project. Restore and build before changing any view model.', 'The existing app still builds.'),
                ('Generate editor properties', 'Use ObservableObject and ObservableProperty', 'Make ItemEditorViewModel partial and inherit ObservableObject. Convert one input field to an observable partial property and build. Convert the remaining inputs and output text properties. The generated property names must match the existing XAML bindings.', 'Typing still updates bound values and the total preview.'),
                ('Keep the editor reactions', 'Use generated change hooks for behavior', 'Implement OnTitleChanged and OnUnitsChanged to update validation. Implement the cost and quantity hooks to update the total and validation. Add NotifyCanExecuteChangedFor to fields affecting Submit.', 'Submit availability and validation respond as each field changes.'),
                ('Generate editor commands', 'Create commands from methods', 'Mark Submit with RelayCommand and CanExecute; mark Cancel with RelayCommand. Remove manual Command construction and ChangeCanExecute calls. Keep the Submitted and CancelRequested events because the page handles navigation.', 'Add, Edit, Cancel, invalid input, and overflow behave as before.'),
                ('Generate BOM commands', 'Apply the same pattern to collection actions', 'Make BillOfMaterialsViewModel partial and inherit ObservableObject. Mark Add, Edit, and Delete methods with RelayCommand. Keep service calls and page requests.', 'Add, Edit, and Delete still update the list.'),
                ('Review generated code', 'Explain what was removed and what remains', 'Compare with Module 7. Point to the generated command names used by XAML. Explain why the validation method is still handwritten.', 'Windows build has zero warnings and errors; the complete workflow works.'),
            ],
            [('ProjectCostEstimator.csproj', source/'ProjectCostEstimator.csproj'), ('ViewModels/ItemEditorViewModel.cs', source/'ViewModels'/'ItemEditorViewModel.cs'), ('ViewModels/BillOfMaterialsViewModel.cs', source/'ViewModels'/'BillOfMaterialsViewModel.cs')],
            ['https://learn.microsoft.com/dotnet/communitytoolkit/mvvm/generators/observableproperty', 'https://learn.microsoft.com/dotnet/communitytoolkit/mvvm/generators/relaycommand'])
    elif number == 9:
        build(9, 'MVVM Navigation and Dialog Boundaries',
            'Complete the dependency injection story by letting the container create both pages and both view models. The BOM view model asks an interface to edit an item or confirm deletion. The MAUI implementation handles navigation and the system alert. This makes the decision logic testable without opening a window.',
            [
                ('Inject the editor view model', 'Use constructor injection for a transient editor', 'Register ItemEditorViewModel and ItemEditorPage as transient services. Change the page constructor to accept the view model. Move edit prefilling into a LoadItem method that the page calls for an existing item.', 'A new editor starts empty and Edit starts with the selected row values.'),
                ('Define interaction results', 'Return a value from a UI workflow', 'Add IBomInteractionService with EditItemAsync and ConfirmDeleteAsync. A null edit result means Cancel or Back; a MaterialItem means Submit.', 'The interface contains no ContentPage, Shell, or DisplayAlert calls.'),
                ('Implement the MAUI adapter', 'Keep navigation and dialogs at the platform boundary', 'Add MauiBomInteractionService. Resolve a fresh editor page from DI, open it with Shell navigation, and await its result when it disappears. Use DisplayAlertAsync for delete confirmation. Unsubscribe the completion handler after use.', 'The editor closes through Submit, Cancel, or Back without adding a row on cancellation.'),
                ('Move the BOM decisions', 'Use async generated commands', 'Inject IBomInteractionService into BillOfMaterialsViewModel. AddAsync adds a returned item. EditAsync replaces only when a result is returned. DeleteAsync removes only after confirmation. The RelayCommand generator creates asynchronous command properties.', 'Add, Edit, Cancel, and both deletion responses work.'),
                ('Reduce page code', 'Keep the page focused on binding', 'Remove BOM page event handlers. Its constructor only accepts the view model and sets BindingContext. Register the interaction implementation in MauiProgram.', 'The page has no BOM data decisions; the view model has no MAUI page dependency.'),
            ],
            [('Services/IBomInteractionService.cs', source/'Services'/'IBomInteractionService.cs'), ('Services/MauiBomInteractionService.cs', source/'Services'/'MauiBomInteractionService.cs'), ('ViewModels/BillOfMaterialsViewModel.cs', source/'ViewModels'/'BillOfMaterialsViewModel.cs'), ('ViewModels/ItemEditorViewModel.cs', source/'ViewModels'/'ItemEditorViewModel.cs'), ('Views/ItemEditorPage.xaml.cs', source/'Views'/'ItemEditorPage.xaml.cs'), ('Views/BillOfMaterialsPage.xaml.cs', source/'Views'/'BillOfMaterialsPage.xaml.cs'), ('MauiProgram.cs', source/'MauiProgram.cs')],
            ['https://learn.microsoft.com/dotnet/maui/fundamentals/dependency-injection/', 'https://learn.microsoft.com/dotnet/communitytoolkit/mvvm/generators/relaycommand'])
    elif number == 10:
        build(10, 'JSON Persistence and View Model Tests',
            'Replace the in-memory BOM service with a JSON-backed implementation while keeping the same view model and editor workflow. Add tests that prove saved items reload and that Cancel or a declined confirmation leaves the collection alone. A visible error message tells the user when loading or saving fails.',
            [
                ('Make storage asynchronous', 'Use Task-returning service operations for file work', 'Change IBomService to LoadAsync, AddItemAsync, ReplaceItemAsync, and RemoveItemAsync. Update the in-memory implementation first so the current app and tests still have a simple substitute.', 'The project builds after the interface and callers are updated.'),
                ('Write a JSON service', 'Separate persisted data from UI behavior', 'Add JsonBomService with a file path constructor. Load the JSON list into its observable collection. Save a temporary file and move it over the previous file after successful serialization. Serialize operations with SemaphoreSlim.', 'A new service instance can reload what a previous instance saved.'),
                ('Handle save failures', 'Preserve the visible collection when storage fails', 'For Add, Edit, and Delete, change the observable collection, attempt to save, and undo the change if the save throws. Let the view model show a concise error.', 'A failed save does not leave an unsaved row on screen.'),
                ('Load when the page appears', 'Connect page lifecycle to view model loading', 'Register JsonBomService as the singleton IBomService using FileSystem.AppDataDirectory. Invoke the generated LoadCommand from the BOM page OnAppearing method. Bind ErrorText above the collection.', 'A new app run restores the saved list, including edits and deletions.'),
                ('Test without a window', 'Replace UI interactions with a fake', 'Create an xUnit test project. Test JSON save, reload, replace, and remove. Use a fake IBomInteractionService to test Cancel, declined Delete, and confirmed Delete through the generated commands.', 'All tests pass, and the Windows MAUI build has zero warnings and errors.'),
            ],
            [('Services/IBomService.cs', source/'Services'/'IBomService.cs'), ('Services/InMemoryBomService.cs', source/'Services'/'InMemoryBomService.cs'), ('Services/JsonBomService.cs', source/'Services'/'JsonBomService.cs'), ('ViewModels/BillOfMaterialsViewModel.cs', source/'ViewModels'/'BillOfMaterialsViewModel.cs'), ('Views/BillOfMaterialsPage.xaml', source/'Views'/'BillOfMaterialsPage.xaml'), ('Views/BillOfMaterialsPage.xaml.cs', source/'Views'/'BillOfMaterialsPage.xaml.cs'), ('MauiProgram.cs', source/'MauiProgram.cs'), ('ProjectCostEstimator.Tests/ProjectCostEstimator.Tests.csproj', ROOT/'ProjectCostEstimator.Tests'/'ProjectCostEstimator.Tests.csproj'), ('ProjectCostEstimator.Tests/BomTests.cs', ROOT/'ProjectCostEstimator.Tests'/'BomTests.cs')],
            ['https://learn.microsoft.com/dotnet/maui/platform-integration/storage/file-system-helpers', 'https://learn.microsoft.com/dotnet/standard/serialization/system-text-json/how-to', 'https://learn.microsoft.com/dotnet/core/testing/unit-testing-csharp-with-xunit'])
    elif number == 11:
        reference = ROOT / 'Module11Reference'
        build(11, 'Rebuild the Estimator with MVVM from the Start',
            'Build the complete IoT cost estimator again from a clean MAUI template. This time begin with the final responsibilities: a model for cost data, a service for stored items, view models for form and list decisions, and views for XAML and platform interactions. Use the reference project only to check your work after each stage. The goal is to make the architecture choices yourself.',
            [
                ('Create a new MAUI project', 'Start the redo from the standard project template', 'In Visual Studio, select Create a new project. Search for and select the .NET MAUI App template, then choose Next. Name the project ProjectCostEstimator, choose a new location outside the earlier module project, and choose Create. Select .NET 10 when Visual Studio asks for a framework. Set the run target to Windows Machine, then build and run the default app before changing any files. Keep Module11Reference closed until a checkpoint requires comparison.', 'The new template app opens with its Hello, World! page and builds with zero errors.'),
                ('Sketch the responsibilities', 'Assign each decision to a layer before coding', 'Draw four boxes: MaterialItem model; IBomService and JsonBomService; ItemEditorViewModel and BillOfMaterialsViewModel; ItemEditorPage and BillOfMaterialsPage. Draw arrows from pages to view models, from the BOM view model to service interfaces, and from the MAUI interaction adapter to pages.', 'You can explain where cost calculation, validation, file IO, navigation, and confirmation belong.'),
                ('Add the model and toolkit', 'Create domain data and generated MVVM plumbing', 'Add the CommunityToolkit.Mvvm package shown in the reference csproj. Create Models/MaterialItem.cs and the small ItemSubmittedEventArgs class. Keep TotalCost derived from cost and quantity. Build before adding pages.', 'A sample item calculates the expected total and the project still builds.'),
                ('Define storage first', 'Let a view model depend on a contract', 'Create Services/IBomService.cs with an observable Items collection and asynchronous Load, Add, Replace, and Remove methods. Add InMemoryBomService as a simple implementation. Use that implementation while developing the pages.', 'The service can add, replace, and remove items without a MAUI window.'),
                ('Build the editor view model', 'Separate form state from Entry controls', 'Create a partial ItemEditorViewModel derived from ObservableObject. Add observable partial properties matching the intended XAML bindings. Add validation, live total, and generated Submit and Cancel commands. Use the change hooks and NotifyCanExecuteChangedFor to update form state.', 'Valid entries enable Submit; blank, negative, invalid, and overflowing values do not.'),
                ('Build the editor page', 'Bind controls and keep navigation at the view boundary', 'Create Views/ItemEditorPage.xaml with Entry bindings, the total and validation Labels, and Submit/Cancel command bindings. Inject the editor view model into the page constructor. The page handles Submitted and CancelRequested and pops navigation. Use LoadItem for editing an existing row.', 'Add mode opens empty; Edit mode is prefilled; Cancel returns no new item.'),
                ('Build the BOM view model', 'Make list commands asynchronous', 'Create BillOfMaterialsViewModel with injected IBomService and IBomInteractionService. Expose the service collection. Implement generated AddAsync, EditAsync, DeleteAsync, and LoadAsync commands. Handle a null editor result as cancellation, and display storage errors through ErrorText.', 'The view model can be tested with fake interactions and no MAUI page.'),
                ('Build the BOM view and UI adapter', 'Keep platform APIs behind a small interface', 'Create the BOM page with CollectionView, row command parameters, an Add button, and an ErrorText label. Create IBomInteractionService and MauiBomInteractionService. The adapter resolves a fresh editor page from DI, awaits its submitted result, and shows delete confirmation.', 'Add, Edit, Cancel, Delete No, and Delete Yes all work with an in-memory BOM.'),
                ('Compose the application', 'Register dependencies in one place', 'Replace AppShell with Home, Bill of Materials, and Help entries. Replace template MainPage with the IoT welcome page under Views and add HelpPage. Register the pages, view models, BOM service, and interaction service in MauiProgram. Copy the finished app artwork only after the workflow runs.', 'Shell navigation opens all pages and the Windows build has zero warnings and errors.'),
                ('Add persistence', 'Change the implementation without changing the BOM workflow', 'Add JsonBomService with a path supplied from FileSystem.AppDataDirectory and register it for IBomService. Save through a temporary file, then replace the previous file. Roll back collection changes when a save fails. Load when the BOM page appears.', 'An added or edited item remains after closing and restarting the app.'),
                ('Test and review', 'Check behavior and boundaries', 'Create a Windows-targeted xUnit project. Test JSON save/reload and the view model decisions with a fake interaction service. Then run the full UI checklist: Add, Edit, Cancel, Delete No, Delete Yes, Help, and restart. Compare your architecture with Module11Reference after trying each stage.', 'Your new project and the reference build, tests pass, and the visible app behaves like the earlier modules.'),
            ],
            [('Services/IBomService.cs', reference/'Services'/'IBomService.cs'), ('Services/JsonBomService.cs', reference/'Services'/'JsonBomService.cs'), ('ViewModels/ItemEditorViewModel.cs', reference/'ViewModels'/'ItemEditorViewModel.cs'), ('ViewModels/BillOfMaterialsViewModel.cs', reference/'ViewModels'/'BillOfMaterialsViewModel.cs'), ('Services/IBomInteractionService.cs', reference/'Services'/'IBomInteractionService.cs'), ('Services/MauiBomInteractionService.cs', reference/'Services'/'MauiBomInteractionService.cs'), ('Views/ItemEditorPage.xaml', reference/'Views'/'ItemEditorPage.xaml'), ('Views/BillOfMaterialsPage.xaml', reference/'Views'/'BillOfMaterialsPage.xaml')],
            ['https://learn.microsoft.com/dotnet/maui/fundamentals/dependency-injection/', 'https://learn.microsoft.com/dotnet/communitytoolkit/mvvm/', 'https://learn.microsoft.com/dotnet/maui/platform-integration/storage/file-system-helpers'],
            starting_point='Create a new .NET MAUI App in Visual Studio with the project name ProjectCostEstimator and the .NET 10 framework. Module11Reference is the complete working result. The earlier module project stays available for comparison. Build the Windows target after each connected group of files, then run its workflow. Write the redo in your new project before comparing with the reference.')
